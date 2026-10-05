#!/usr/bin/env python3
"""Gera o template oficial de importação de vendas do Intellix (.xlsx).

O arquivo resultante é o que o cliente baixa, preenche e sobe na tela de
importação. Os cabeçalhos batem exatamente com o que o importador do backend
espera (ver backend/app/importacao/leitor_excel.py). Rodar:

    pip install openpyxl
    python scripts/gerar_template.py

Saída: frontend/public/template_intellix_vendas.xlsx
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (coluna, rótulo amigável, obrigatória?, exemplo1, exemplo2)
COLUNAS = [
    ("venda_id", "ID da venda", False, "V001", "V001"),
    ("data_venda", "Data da venda", True, "2026-03-05", "2026-03-05"),
    ("produto", "Produto", True, "Tênis casual", "Boné"),
    ("categoria", "Categoria", False, "Calçados", "Acessórios"),
    ("cliente", "Cliente", False, "Ana Silva", "Ana Silva"),
    ("vendedor", "Vendedor", False, "Carlos Nunes", "Carlos Nunes"),
    ("canal", "Canal", False, "loja", "loja"),
    ("forma_pagamento", "Forma de pagamento", False, "pix", "pix"),
    ("quantidade", "Quantidade", True, 1, 2),
    ("valor_unitario", "Valor unitário", False, 279.90, 59.90),
    ("desconto", "Desconto", False, 0, 0),
    ("valor_total", "Valor total", False, 279.90, 119.80),
]

AZUL = "1F5C93"
CREME = "FBF9F3"


def gerar(destino: str) -> None:
    wb = Workbook()

    # ---- Aba de dados ----
    ws = wb.active
    ws.title = "Vendas"

    cabecalho_fill = PatternFill("solid", fgColor=AZUL)
    cabecalho_font = Font(bold=True, color="FFFFFF")

    for i, (coluna, _rotulo, obrig, _e1, _e2) in enumerate(COLUNAS, start=1):
        titulo = coluna + (" *" if obrig else "")
        c = ws.cell(row=1, column=i, value=titulo)
        c.fill = cabecalho_fill
        c.font = cabecalho_font
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(coluna) + 4)

    for linha, idx in ((2, 3), (3, 4)):
        for i, col in enumerate(COLUNAS, start=1):
            ws.cell(row=linha, column=i, value=col[idx])

    ws.freeze_panes = "A2"

    # ---- Aba de instruções ----
    ins = wb.create_sheet("Instruções")
    ins.column_dimensions["A"].width = 26
    ins.column_dimensions["B"].width = 80

    def linha_ins(a, b, negrito=False):
        ra = ins.cell(row=linha_ins.n, column=1, value=a)
        rb = ins.cell(row=linha_ins.n, column=2, value=b)
        if negrito:
            ra.font = Font(bold=True)
            rb.font = Font(bold=True)
        rb.alignment = Alignment(wrap_text=True, vertical="top")
        linha_ins.n += 1

    linha_ins.n = 1
    linha_ins("Template Intellix", "Planilha de vendas para importação", negrito=True)
    linha_ins("", "")
    linha_ins("Como usar", "Preencha a aba \"Vendas\", uma linha por item vendido. "
                            "Apague as duas linhas de exemplo antes de enviar.")
    linha_ins("", "")
    linha_ins("Colunas obrigatórias", "data_venda, produto, quantidade e pelo menos "
                                      "um entre valor_unitario e valor_total.")
    linha_ins("Marcação *", "O asterisco no cabeçalho indica coluna obrigatória. "
                            "Não renomeie os cabeçalhos.")
    linha_ins("", "")
    linha_ins("data_venda", "Data no formato AAAA-MM-DD (ex.: 2026-03-05) ou DD/MM/AAAA. "
                            "Entre 2023 e 2026.")
    linha_ins("quantidade", "Número inteiro de unidades (ex.: 2).")
    linha_ins("valor_unitario", "Preço de uma unidade (ex.: 279.90). Se informar só o "
                                "valor_total, deixe em branco.")
    linha_ins("valor_total", "Total do item (quantidade x unitário - desconto). Se "
                             "deixar em branco, calculamos a partir do unitário.")
    linha_ins("venda_id", "Opcional: agrupa itens da mesma venda (ex.: V001 em duas "
                          "linhas = uma venda com dois itens). Se deixar em branco, "
                          "cada linha vira uma venda.")
    linha_ins("canal", "Ex.: loja, site, whatsapp, marketplace.")
    linha_ins("forma_pagamento", "Ex.: pix, cartao_credito, cartao_debito, dinheiro, boleto.")
    linha_ins("", "")
    linha_ins("Colunas extras", "Qualquer coluna fora desta lista é ignorada na importação.")

    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    print(f"Template gerado em: {destino}")


if __name__ == "__main__":
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    gerar(os.path.join(raiz, "frontend", "public", "template_intellix_vendas.xlsx"))
