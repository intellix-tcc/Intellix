"""Leitura e validação de planilhas de venda no servidor.

Espelha deliberadamente a validação que o frontend já faz em
`frontend/src/servicos/importacao.js` (mesmas colunas, mesmos sinônimos):
o navegador barra o óbvio antes de gastar upload, e aqui a gente confere de
novo porque o servidor nunca confia no cliente. As duas listas precisam
dizer a mesma coisa -- se mudar uma, mude a outra.

Só `.xlsx` (openpyxl não lê o `.xls` binário antigo). O leitor devolve as
linhas já mapeadas para as colunas canônicas de `staging_venda`; a conversão
de tipo e a carga no star schema ficam no `etl.py`.
"""

from __future__ import annotations

import io
import unicodedata
from typing import Any

TAMANHO_MAX_MB = 10

# Colunas da staging_venda (001_schema.sql). A planilha pode vir com acento,
# maiúscula ou espaço -- normalizamos antes de comparar.
COLUNAS_CONHECIDAS = [
    "data_venda",
    "produto",
    "categoria",
    "cliente",
    "vendedor",
    "canal",
    "forma_pagamento",
    "quantidade",
    "valor_unitario",
    "desconto",
    "valor_total",
    "venda_id",
]

COLUNAS_OBRIGATORIAS = ["data_venda", "produto", "quantidade"]
COLUNAS_VALOR = ["valor_unitario", "valor_total"]  # pelo menos uma

# Sinônimos aceitos (mesmo conjunto do frontend + "venda_id").
SINONIMOS = {
    "data_venda": ["data", "data da venda", "dt venda", "emissao", "data emissao"],
    "produto": ["item", "descricao", "descricao do produto", "mercadoria"],
    "categoria": ["grupo", "linha", "departamento"],
    "cliente": ["comprador", "nome do cliente"],
    "vendedor": ["atendente", "responsavel"],
    "canal": ["origem", "canal de venda"],
    "forma_pagamento": ["pagamento", "forma de pagamento", "meio de pagamento"],
    "quantidade": ["qtd", "qtde", "quant"],
    "valor_unitario": ["preco", "preco unitario", "valor unit", "unitario"],
    "desconto": ["descontos", "abatimento"],
    "valor_total": ["total", "valor", "valor da venda", "faturamento"],
    "venda_id": ["id venda", "id da venda", "pedido", "nota", "cupom", "numero da venda"],
}

# .xlsx é um ZIP -> começa com "PK\x03\x04".
_ASSINATURA_ZIP = b"PK\x03\x04"


class ErroImportacao(Exception):
    """Erro com mensagem pronta para o usuário (não vaza stack trace)."""

    def __init__(self, mensagem: str):
        super().__init__(mensagem)
        self.mensagem = mensagem


def normalizar(texto: Any) -> str:
    s = "" if texto is None else str(texto)
    s = s.strip().lower()
    # tira acentos (NFD + remove marcas de combinação)
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    out = []
    for ch in s:
        if ch.isalnum() or ch == "_":
            out.append(ch)
        elif ch in (" ", "-"):
            out.append("_")
    # colapsa underscores repetidos
    resultado = "".join(out)
    while "__" in resultado:
        resultado = resultado.replace("__", "_")
    return resultado.strip("_")


def mapear_coluna(cabecalho: Any) -> str | None:
    """Cabeçalho da planilha -> coluna canônica, ou None se não bater."""
    alvo = normalizar(cabecalho)
    if not alvo:
        return None
    if alvo in COLUNAS_CONHECIDAS:
        return alvo
    for coluna, apelidos in SINONIMOS.items():
        if any(normalizar(a) == alvo for a in apelidos):
            return coluna
    return None


def _assinatura_valida(conteudo: bytes) -> bool:
    return conteudo[:4] == _ASSINATURA_ZIP


def ler_planilha(conteudo: bytes, nome_arquivo: str = "") -> dict:
    """Abre o .xlsx e devolve:

        {
          "aba": str,
          "cabecalho": list[str],
          "colunas_encontradas": list[str],   # canônicas
          "colunas_faltando": list[str],
          "extras": list[str],                 # ignoradas
          "linhas": list[dict],                # {coluna_canonica: valor, "_linha": n}
          "total_linhas": int,
        }

    Levanta ErroImportacao (mensagem de usuário) quando não dá para ler.
    """
    nome = (nome_arquivo or "").lower()
    if nome.endswith(".xls"):
        raise ErroImportacao(
            "O formato .xls antigo não é aceito no envio. Abra no Excel e salve "
            "como .xlsx."
        )
    if len(conteudo) == 0:
        raise ErroImportacao("Esse arquivo está vazio.")
    if len(conteudo) > TAMANHO_MAX_MB * 1024 * 1024:
        raise ErroImportacao(f"A planilha precisa ter no máximo {TAMANHO_MAX_MB} MB.")
    if not _assinatura_valida(conteudo):
        raise ErroImportacao(
            "Não conseguimos processar esse arquivo. Verifique se ele está em "
            "formato Excel (.xlsx) válido e se não está protegido por senha."
        )

    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 -- qualquer falha de leitura vira mensagem de usuário
        raise ErroImportacao(
            "Não conseguimos abrir essa planilha. Verifique se o arquivo é um "
            "Excel válido e se não está protegido por senha."
        )

    if not wb.sheetnames:
        raise ErroImportacao("Essa planilha não tem nenhuma aba.")

    aba = wb.sheetnames[0]
    ws = wb[aba]

    linhas_brutas = list(ws.iter_rows(values_only=True))
    wb.close()

    # descarta linhas totalmente vazias do fim/meio
    def _vazia(linha) -> bool:
        return all(c is None or str(c).strip() == "" for c in linha)

    linhas_brutas = [linha for linha in linhas_brutas if not _vazia(linha)]

    if len(linhas_brutas) == 0:
        raise ErroImportacao(
            f'A aba "{aba}" está vazia. Salve a planilha com os dados e tente de novo.'
        )
    if len(linhas_brutas) == 1:
        raise ErroImportacao(
            f'A aba "{aba}" só tem o cabeçalho, sem nenhuma linha de venda.'
        )

    cabecalho = list(linhas_brutas[0])
    mapa = [mapear_coluna(c) for c in cabecalho]
    encontradas = [c for c in mapa if c]

    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in encontradas]
    if not any(c in encontradas for c in COLUNAS_VALOR):
        faltando.append("valor_total")

    extras = [str(c) for c, m in zip(cabecalho, mapa) if c and not m]

    if faltando:
        raise ErroImportacao(
            "Faltam colunas obrigatórias na planilha: "
            + ", ".join(faltando)
            + ". Use o template oficial ou renomeie os cabeçalhos."
        )

    linhas: list[dict] = []
    # linha_origem conta a partir de 1 = cabeçalho; dados começam em 2.
    for n, linha in enumerate(linhas_brutas[1:], start=2):
        registro: dict[str, Any] = {"_linha": n}
        for idx, coluna in enumerate(mapa):
            if coluna and idx < len(linha):
                registro[coluna] = linha[idx]
        linhas.append(registro)

    return {
        "aba": aba,
        "cabecalho": [str(c) for c in cabecalho],
        "colunas_encontradas": encontradas,
        "colunas_faltando": [],
        "extras": extras,
        "linhas": linhas,
        "total_linhas": len(linhas),
    }
